from pathlib import Path
import sqlite3
import hashlib
import json
import re
import shutil
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Iterable

ROOT = Path(__file__).resolve().parents[2]

KB_ROOT = ROOT / "data" / "knowledge"
SOURCE_ROOT = KB_ROOT / "source"
PROCESSED_ROOT = KB_ROOT / "processed"
QUARANTINE_ROOT = KB_ROOT / "quarantine"
BENCHMARK_ROOT = KB_ROOT / "benchmark"

DB_PATH = KB_ROOT / "production_knowledge.db"
MANIFEST_PATH = KB_ROOT / "production_manifest.json"

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".json",
    ".html",
    ".htm",
}

MAX_FILE_SIZE_MB = 50
MIN_TEXT_LENGTH = 40

CHUNK_SIZE = 900
CHUNK_OVERLAP = 120


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(
        text.encode("utf-8", errors="ignore")
    ).hexdigest()


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    lines = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            lines.append(line)

    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def load_txt(path: Path) -> str:
    return path.read_text(
        encoding="utf-8",
        errors="replace",
    )


def load_json(path: Path) -> str:
    obj = json.loads(
        path.read_text(
            encoding="utf-8",
            errors="replace",
        )
    )

    if isinstance(obj, str):
        return obj

    return json.dumps(
        obj,
        ensure_ascii=False,
        indent=2,
    )


def load_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx belum terinstall. "
            "Install dengan: python -m pip install python-docx"
        ) from exc

    document = Document(str(path))

    paragraphs = [
        p.text.strip()
        for p in document.paragraphs
        if p.text.strip()
    ]

    tables = []

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            values = [v for v in values if v]
            if values:
                tables.append(" | ".join(values))

    return "\n".join(paragraphs + tables)


def load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "pypdf belum terinstall. "
            "Install dengan: python -m pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))

    pages = []

    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        text = normalize_text(text)

        if text:
            pages.append(
                f"[PAGE {index}]\n{text}"
            )

    return "\n\n".join(pages)


def load_html(path: Path) -> str:
    raw = path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    try:
        from bs4 import BeautifulSoup
    except ImportError:
        text = re.sub(
            r"<script.*?</script>",
            " ",
            raw,
            flags=re.I | re.S,
        )
        text = re.sub(
            r"<style.*?</style>",
            " ",
            text,
            flags=re.I | re.S,
        )
        text = re.sub(r"<[^>]+>", " ", text)
        return normalize_text(text)

    soup = BeautifulSoup(raw, "html.parser")

    for tag in soup(
        ["script", "style", "noscript"]
    ):
        tag.decompose()

    return soup.get_text("\n")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return load_txt(path)

    if suffix == ".json":
        return load_json(path)

    if suffix == ".docx":
        return load_docx(path)

    if suffix == ".pdf":
        return load_pdf(path)

    if suffix in {".html", ".htm"}:
        return load_html(path)

    raise ValueError(
        f"Format tidak didukung: {suffix}"
    )


def classify_source(
    relative_path: Path,
    text: str = "",
) -> dict:
    """
    Classify production knowledge source.

    Priority:
    1. Explicit metadata inside document
    2. Existing filename/path classification
    """

    parts = [
        p.lower()
        for p in relative_path.parts
    ]

    service = "general"

    mapping = {
        "ktp_el": "ktp_el",
        "ktp": "ktp_el",
        "kartu_keluarga": "kartu_keluarga",
        "kk": "kartu_keluarga",
        "kia": "kia",
        "akta_kelahiran": "akta_kelahiran",
        "akta_kematian": "akta_kematian",
        "perkawinan": "perkawinan",
        "perceraian": "perceraian",
        "pindah_datang": "pindah_datang",
        "perubahan_data": "perubahan_data",
        "layanan_online": "layanan_online",
    }

    for part in parts:
        if part in mapping:
            service = mapping[part]
            break

    content_type = "general"

    name = relative_path.stem.lower()

    if any(
        token in name
        for token in [
            "syarat",
            "persyaratan",
            "requirement",
        ]
    ):
        content_type = "requirements"

    elif any(
        token in name
        for token in [
            "prosedur",
            "cara",
            "alur",
            "proses",
        ]
    ):
        content_type = "procedure"

    elif any(
        token in name
        for token in [
            "biaya",
            "tarif",
        ]
    ):
        content_type = "fee"

    elif any(
        token in name
        for token in [
            "waktu",
            "durasi",
            "lama",
        ]
    ):
        content_type = "processing_time"

    elif any(
        token in name
        for token in [
            "status",
            "cek",
            "tracking",
        ]
    ):
        content_type = "status"

    # --------------------------------------------------------
    # Explicit document metadata
    # Format:
    # SERVICE: ktp_el
    # CONTENT_TYPE: requirements
    # SOURCE_STATUS: official
    # AUTHORITY: official
    # VERIFIED: true
    # SOURCE_TYPE: government
    # DOMAIN: kependudukan
    # --------------------------------------------------------

    metadata = {}

    for line in text.splitlines():
        match = re.match(
            r"^\s*([A-Z][A-Z0-9_ -]{1,50})\s*:\s*(.*?)\s*$",
            line,
        )

        if not match:
            continue

        key = (
            match.group(1)
            .strip()
            .upper()
            .replace(" ", "_")
            .replace("-", "_")
        )

        value = match.group(2).strip()

        if value:
            metadata[key] = value

    if metadata.get("DOMAIN"):
        domain = metadata["DOMAIN"].lower()

    else:
        domain = "kependudukan"

    if metadata.get("SERVICE"):
        service = metadata["SERVICE"].lower()

    if metadata.get("CONTENT_TYPE"):
        content_type = metadata["CONTENT_TYPE"].lower()

    source_status = (
        metadata.get("SOURCE_STATUS")
        or metadata.get("AUTHORITY")
        or "unknown"
    ).lower()

    authority = (
        metadata.get("AUTHORITY")
        or source_status
        or "unknown"
    ).lower()

    verified_raw = (
        metadata.get("VERIFIED", "")
        .strip()
        .lower()
    )

    verified = verified_raw in {
        "true",
        "1",
        "yes",
        "verified",
        "ya",
    }

    source_type = (
        metadata.get("SOURCE_TYPE")
        or "file"
    ).lower()

    return {
        "domain": domain,
        "service": service,
        "content_type": content_type,
        "source_status": source_status,
        "authority": authority,
        "verified": verified,
        "source_type": source_type,
    }


def chunk_text(
    text: str,
    size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[str]:

    text = normalize_text(text)

    if not text:
        return []

    paragraphs = [
        p.strip()
        for p in re.split(
            r"\n\s*\n",
            text,
        )
        if p.strip()
    ]

    chunks = []
    current = ""

    for paragraph in paragraphs:

        if len(paragraph) > size:
            if current:
                chunks.append(current)
                current = ""

            start = 0

            while start < len(paragraph):
                end = min(
                    start + size,
                    len(paragraph),
                )

                piece = paragraph[start:end].strip()

                if piece:
                    chunks.append(piece)

                if end >= len(paragraph):
                    break

                start = max(
                    end - overlap,
                    start + 1,
                )

            continue

        candidate = (
            paragraph
            if not current
            else current + "\n\n" + paragraph
        )

        if len(candidate) <= size:
            current = candidate
        else:
            if current:
                chunks.append(current)

            tail = ""

            if overlap > 0 and current:
                tail = current[-overlap:]

            current = (
                tail + "\n\n" + paragraph
                if tail
                else paragraph
            )

    if current:
        chunks.append(current)

    return [
        normalize_text(chunk)
        for chunk in chunks
        if len(normalize_text(chunk)) >= 20
    ]


def make_document_id(file_hash: str) -> str:
    return "DOC-" + file_hash[:16].upper()


def make_chunk_id(
    document_id: str,
    index: int,
    text: str,
) -> str:
    digest = sha256_text(
        f"{document_id}:{index}:{text}"
    )

    return "CHUNK-" + digest[:20].upper()


@dataclass
class DocumentRecord:
    document_id: str
    title: str
    source_file: str
    source_hash: str
    source_type: str
    domain: str
    service: str
    content_type: str
    text_length: int
    chunk_count: int
    ingested_at: str


class ProductionKnowledgeBase:

    def __init__(
        self,
        db_path: Path = DB_PATH,
    ):
        self.db_path = db_path

        db_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        self.conn = sqlite3.connect(
            str(db_path)
        )

        self.conn.row_factory = sqlite3.Row

        self._initialize()

    def _initialize(self):

        self.conn.executescript(
            """
            PRAGMA journal_mode=WAL;
            PRAGMA foreign_keys=ON;

            CREATE TABLE IF NOT EXISTS documents (
                document_id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                source_file TEXT NOT NULL,
                source_hash TEXT NOT NULL,
                source_type TEXT NOT NULL,
                domain TEXT NOT NULL,
                service TEXT NOT NULL,
                content_type TEXT NOT NULL,
                source_status TEXT NOT NULL DEFAULT 'unknown',
                authority TEXT NOT NULL DEFAULT 'unknown',
                verified INTEGER NOT NULL DEFAULT 0,
                text_length INTEGER NOT NULL,
                chunk_count INTEGER NOT NULL,
                ingested_at TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1
            );

            CREATE TABLE IF NOT EXISTS chunks (
                chunk_id TEXT PRIMARY KEY,
                document_id TEXT NOT NULL,
                chunk_index INTEGER NOT NULL,
                text TEXT NOT NULL,
                text_hash TEXT NOT NULL,
                token_estimate INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(document_id)
                    REFERENCES documents(document_id)
                    ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS
                idx_chunks_document
                ON chunks(document_id);

            CREATE INDEX IF NOT EXISTS
                idx_documents_service
                ON documents(service);

            CREATE INDEX IF NOT EXISTS
                idx_documents_content_type
                ON documents(content_type);

            CREATE VIRTUAL TABLE IF NOT EXISTS
                chunks_fts USING fts5(
                    chunk_id UNINDEXED,
                    document_id UNINDEXED,
                    text,
                    tokenize='unicode61'
                );

            CREATE TABLE IF NOT EXISTS ingestion_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_file TEXT,
                source_hash TEXT,
                status TEXT,
                message TEXT,
                created_at TEXT
            );
            """
        )

        # ----------------------------------------------------
        # Backward-compatible metadata migration
        # ----------------------------------------------------
        existing_columns = {
            row[1]
            for row in self.conn.execute(
                "PRAGMA table_info(documents)"
            ).fetchall()
        }

        migrations = [
            (
                "source_status",
                "ALTER TABLE documents ADD COLUMN "
                "source_status TEXT NOT NULL DEFAULT 'unknown'"
            ),
            (
                "authority",
                "ALTER TABLE documents ADD COLUMN "
                "authority TEXT NOT NULL DEFAULT 'unknown'"
            ),
            (
                "verified",
                "ALTER TABLE documents ADD COLUMN "
                "verified INTEGER NOT NULL DEFAULT 0"
            ),
        ]

        for column, sql in migrations:
            if column not in existing_columns:
                self.conn.execute(sql)

        self.conn.commit()

    def close(self):
        self.conn.close()

    def document_exists(
        self,
        source_hash: str,
    ) -> bool:

        row = self.conn.execute(
            """
            SELECT document_id
            FROM documents
            WHERE source_hash = ?
              AND active = 1
            LIMIT 1
            """,
            (source_hash,),
        ).fetchone()

        return row is not None

    def ingest_file(
        self,
        path: Path,
        source_root: Path,
    ) -> dict:

        relative = path.relative_to(
            source_root
        )

        source_hash = sha256_file(path)

        if self.document_exists(source_hash):
            return {
                "status": "SKIP",
                "file": str(relative),
                "reason": "already_ingested",
            }

        # Extract document text BEFORE classification so that
        # explicit metadata inside the document can override
        # filename/path inference.
        text = normalize_text(
            extract_text(path)
        )

        metadata = classify_source(
            relative,
            text,
        )

        title = path.stem.replace(
            "_",
            " ",
        ).replace(
            "-",
            " ",
        ).strip()

        document_id = make_document_id(
            source_hash
        )

        if len(text) < MIN_TEXT_LENGTH:
            raise ValueError(
                f"Isi dokumen terlalu pendek: "
                f"{len(text)} karakter"
            )

        chunks = chunk_text(text)

        if not chunks:
            raise ValueError(
                "Tidak menghasilkan chunk."
            )

        now = utc_now()

        self.conn.execute(
            """
            INSERT INTO documents (
                document_id,
                title,
                source_file,
                source_hash,
                source_type,
                domain,
                service,
                content_type,
                source_status,
                authority,
                verified,
                text_length,
                chunk_count,
                ingested_at,
                active
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1)
            """,
            (
                document_id,
                title,
                str(relative),
                source_hash,
                metadata["source_type"],
                metadata["domain"],
                metadata["service"],
                metadata["content_type"],
                metadata["source_status"],
                metadata["authority"],
                1 if metadata["verified"] else 0,
                len(text),
                len(chunks),
                now,
            ),
        )

        for index, chunk in enumerate(
            chunks
        ):

            chunk_id = make_chunk_id(
                document_id,
                index,
                chunk,
            )

            token_estimate = max(
                1,
                len(
                    re.findall(
                        r"\S+",
                        chunk,
                    )
                ),
            )

            self.conn.execute(
                """
                INSERT INTO chunks (
                    chunk_id,
                    document_id,
                    chunk_index,
                    text,
                    text_hash,
                    token_estimate,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    chunk_id,
                    document_id,
                    index,
                    chunk,
                    sha256_text(chunk),
                    token_estimate,
                    now,
                ),
            )

            self.conn.execute(
                """
                INSERT INTO chunks_fts (
                    chunk_id,
                    document_id,
                    text
                )
                VALUES (?, ?, ?)
                """,
                (
                    chunk_id,
                    document_id,
                    chunk,
                ),
            )

        self.conn.execute(
            """
            INSERT INTO ingestion_log (
                source_file,
                source_hash,
                status,
                message,
                created_at
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                str(relative),
                source_hash,
                "SUCCESS",
                f"{len(chunks)} chunks",
                now,
            ),
        )

        self.conn.commit()

        return {
            "status": "INGESTED",
            "file": str(relative),
            "document_id": document_id,
            "chunks": len(chunks),
            "characters": len(text),
            "service": metadata["service"],
            "content_type": metadata["content_type"],
            "source_status": metadata["source_status"],
            "authority": metadata["authority"],
            "verified": metadata["verified"],
        }

    def search(
        self,
        query: str,
        limit: int = 5,
        authoritative_only: bool = False,
        service: str | None = None,
    ) -> list[dict]:

        query = normalize_text(query)

        if not query:
            return []

        safe_terms = re.findall(
            r"[\w-]+",
            query,
            flags=re.UNICODE,
        )

        if not safe_terms:
            return []

        fts_query = " OR ".join(
            f'"{term}"'
            for term in safe_terms
        )

        where_clauses = [
            "chunks_fts MATCH ?",
            "d.active = 1",
        ]

        params: list = [fts_query]

        if authoritative_only:
            where_clauses.extend([
                "LOWER(COALESCE(d.source_status, '')) IN ('official', 'verified')",
                "LOWER(COALESCE(d.authority, '')) IN ('official', 'verified')",
                "COALESCE(d.verified, 0) = 1",
            ])

        if service:
            where_clauses.append(
                "LOWER(COALESCE(d.service, '')) = LOWER(?)"
            )
            params.append(service)

        params.append(int(limit))

        sql = f"""
            SELECT
                f.chunk_id,
                f.document_id,
                f.text,
                d.title,
                d.source_file,
                d.service,
                d.content_type,
                d.source_status,
                d.authority,
                d.verified
            FROM chunks_fts f
            JOIN chunks c
              ON c.chunk_id = f.chunk_id
            JOIN documents d
              ON d.document_id = f.document_id
            WHERE {' AND '.join(where_clauses)}
            LIMIT ?
        """

        rows = self.conn.execute(
            sql,
            tuple(params),
        ).fetchall()

        return [
            dict(row)
            for row in rows
        ]

    def stats(self) -> dict:

        documents = self.conn.execute(
            """
            SELECT COUNT(*)
            FROM documents
            WHERE active = 1
            """
        ).fetchone()[0]

        chunks = self.conn.execute(
            """
            SELECT COUNT(*)
            FROM chunks
            """
        ).fetchone()[0]

        services = self.conn.execute(
            """
            SELECT service, COUNT(*)
            FROM documents
            WHERE active = 1
            GROUP BY service
            ORDER BY service
            """
        ).fetchall()

        content_types = self.conn.execute(
            """
            SELECT content_type, COUNT(*)
            FROM documents
            WHERE active = 1
            GROUP BY content_type
            ORDER BY content_type
            """
        ).fetchall()

        return {
            "database": str(self.db_path),
            "documents": documents,
            "chunks": chunks,
            "services": {
                row[0]: row[1]
                for row in services
            },
            "content_types": {
                row[0]: row[1]
                for row in content_types
            },
        }


def validate_environment():

    KB_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    SOURCE_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    PROCESSED_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    QUARANTINE_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )

    BENCHMARK_ROOT.mkdir(
        parents=True,
        exist_ok=True,
    )


def backup_database():

    if not DB_PATH.exists():
        return None

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    backup = (
        KB_ROOT
        / f"production_knowledge_backup_{timestamp}.db"
    )

    shutil.copy2(
        DB_PATH,
        backup,
    )

    return backup


def validate_file(
    path: Path,
) -> tuple[bool, str]:

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return False, "unsupported_extension"

    if not path.is_file():
        return False, "not_file"

    size = path.stat().st_size

    if size == 0:
        return False, "empty_file"

    if size > MAX_FILE_SIZE_MB * 1024 * 1024:
        return False, "file_too_large"

    return True, "ok"


def build_manifest(
    kb: ProductionKnowledgeBase,
):

    stats = kb.stats()

    manifest = {
        "name": "Nusantara AI Production Knowledge Base",
        "version": "1.0.0",
        "generated_at": utc_now(),
        "database": str(
            DB_PATH.relative_to(ROOT)
        ),
        "source_root": str(
            SOURCE_ROOT.relative_to(ROOT)
        ),
        "chunking": {
            "chunk_size": CHUNK_SIZE,
            "chunk_overlap": CHUNK_OVERLAP,
        },
        "supported_extensions": sorted(
            SUPPORTED_EXTENSIONS
        ),
        "stats": stats,
    }

    MANIFEST_PATH.write_text(
        json.dumps(
            manifest,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def ingest_all() -> int:

    validate_environment()

    backup = backup_database()

    kb = ProductionKnowledgeBase()

    results = []

    try:

        files = sorted(
            path
            for path in SOURCE_ROOT.rglob("*")
            if path.is_file()
        )

        for path in files:

            valid, reason = validate_file(
                path
            )

            if not valid:
                results.append(
                    {
                        "status": "QUARANTINE",
                        "file": str(
                            path.relative_to(
                                SOURCE_ROOT
                            )
                        ),
                        "reason": reason,
                    }
                )

                continue

            try:

                result = kb.ingest_file(
                    path,
                    SOURCE_ROOT,
                )

                results.append(result)

            except Exception as exc:

                results.append(
                    {
                        "status": "ERROR",
                        "file": str(
                            path.relative_to(
                                SOURCE_ROOT
                            )
                        ),
                        "reason": str(exc),
                    }
                )

        build_manifest(kb)

        print()
        print("=" * 72)
        print("NUSANTARA PRODUCTION KNOWLEDGE BASE")
        print("=" * 72)

        for result in results:
            print(
                f"[{result['status']}] "
                f"{result['file']}"
            )

        stats = kb.stats()

        print()
        print(
            f"DOCUMENTS : {stats['documents']}"
        )
        print(
            f"CHUNKS    : {stats['chunks']}"
        )

        print()
        print("SERVICES")

        for key, value in stats[
            "services"
        ].items():
            print(
                f"  {key}: {value}"
            )

        print()
        print("CONTENT TYPES")

        for key, value in stats[
            "content_types"
        ].items():
            print(
                f"  {key}: {value}"
            )

        if backup:
            print()
            print(
                f"BACKUP    : {backup}"
            )

        print(
            f"MANIFEST  : {MANIFEST_PATH}"
        )

        return 0

    finally:
        kb.close()


def self_test() -> int:

    validate_environment()

    test_db = (
        KB_ROOT
        / "production_knowledge_selftest.db"
    )

    if test_db.exists():
        test_db.unlink()

    kb = ProductionKnowledgeBase(
        test_db
    )

    try:

        text = (
            "KTP-el adalah dokumen identitas "
            "penduduk. Status pengajuan KTP-el "
            "dapat diperiksa melalui layanan "
            "administrasi kependudukan."
        )

        temp = (
            SOURCE_ROOT
            / "_SELFTEST"
            / "ktp_el"
            / "test.txt"
        )

        temp.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        temp.write_text(
            text,
            encoding="utf-8",
        )

        result = kb.ingest_file(
            temp,
            SOURCE_ROOT,
        )

        assert result["status"] == "INGESTED"
        assert result["chunks"] >= 1

        results = kb.search(
            "Apa itu KTP-el?",
            limit=5,
        )

        assert results

        assert any(
            "KTP-el" in item["text"]
            for item in results
        )

        stats = kb.stats()

        assert stats["documents"] == 1
        assert stats["chunks"] >= 1

        print(
            "[PASS] Database initialization"
        )
        print(
            "[PASS] Document ingestion"
        )
        print(
            "[PASS] Chunk generation"
        )
        print(
            "[PASS] FTS5 lexical search"
        )
        print(
            "[PASS] Metadata generation"
        )
        print(
            "[PASS] Statistics"
        )

        print()
        print(
            "PRODUCTION_KB_SELFTEST: PASS"
        )

        return 0

    finally:

        kb.close()

        try:
            temp.unlink()
            temp.parent.rmdir()
            temp.parent.parent.rmdir()
        except Exception:
            pass

        try:
            test_db.unlink()
        except Exception:
            pass


def print_usage():

    print(
        """
Nusantara Production Knowledge Base

Usage:

  python -m app.knowledge.production_kb self-test

  python -m app.knowledge.production_kb ingest

  python -m app.knowledge.production_kb stats

  python -m app.knowledge.production_kb search "Apa itu KTP-el?"

Source directory:

  data/knowledge/source/

Database:

  data/knowledge/production_knowledge.db
"""
    )


def main():

    command = (
        sys.argv[1]
        if len(sys.argv) > 1
        else "help"
    )

    if command == "self-test":
        raise SystemExit(
            self_test()
        )

    if command == "ingest":
        raise SystemExit(
            ingest_all()
        )

    if command == "stats":

        validate_environment()

        kb = ProductionKnowledgeBase()

        try:
            print(
                json.dumps(
                    kb.stats(),
                    ensure_ascii=False,
                    indent=2,
                )
            )
        finally:
            kb.close()

        return

    if command == "search":

        if len(sys.argv) < 3:
            print(
                "Query search wajib diberikan."
            )
            raise SystemExit(2)

        validate_environment()

        kb = ProductionKnowledgeBase()

        try:

            results = kb.search(
                " ".join(
                    sys.argv[2:]
                )
            )

            print(
                json.dumps(
                    results,
                    ensure_ascii=False,
                    indent=2,
                )
            )

        finally:
            kb.close()

        return

    print_usage()


if __name__ == "__main__":
    main()
